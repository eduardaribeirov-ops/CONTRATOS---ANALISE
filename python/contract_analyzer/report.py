"""Geração do relatório de análise em .docx."""

from __future__ import annotations

import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

GRAVIDADE_CORES = {
    "Alta": RGBColor(0xC0, 0x00, 0x00),
    "Média": RGBColor(0xB8, 0x86, 0x0B),
    "Baixa": RGBColor(0x1E, 0x7B, 0x34),
}


def build_report(analysis: dict, source_filename: str, output_path: str) -> None:
    document = Document()

    title = document.add_heading("Análise de Contrato", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"Arquivo analisado: {source_filename}\n"
        f"Data da análise: {datetime.date.today().strftime('%d/%m/%Y')}"
    )
    meta_run.italic = True

    document.add_paragraph()

    resumo = analysis.get("resumo", {})
    document.add_heading("Resumo do Contrato", level=1)
    _add_field(document, "Partes", resumo.get("partes"))
    _add_field(document, "Objeto", resumo.get("objeto"))
    _add_field(document, "Valor e pagamento", resumo.get("valor_e_pagamento"))
    _add_field(document, "Prazo e vigência", resumo.get("prazo_vigencia"))

    pontos = analysis.get("pontos", [])
    document.add_heading(f"Pontos de Melhoria / Ajuste Necessários ({len(pontos)})", level=1)

    if not pontos:
        document.add_paragraph("Nenhum ponto relevante identificado.")
    else:
        for idx, ponto in enumerate(pontos, start=1):
            _add_ponto(document, idx, ponto)

    document.add_heading("Conclusão", level=1)
    document.add_paragraph(analysis.get("conclusao", ""))

    document.add_paragraph()
    aviso = document.add_paragraph()
    aviso_run = aviso.add_run(
        "Este documento é gerado automaticamente por um checklist de regras "
        "(busca por palavras-chave, sem uso de IA) e tem caráter apenas "
        "informativo/preparatório - pode haver falsos positivos e falsos "
        "negativos. Recomenda-se revisão por um advogado antes de qualquer decisão."
    )
    aviso_run.italic = True
    aviso_run.font.size = Pt(9)

    document.save(output_path)


def _add_field(document: Document, label: str, value: str | None) -> None:
    p = document.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value or "não identificado")


def _add_ponto(document: Document, idx: int, ponto: dict) -> None:
    gravidade = ponto.get("gravidade", "")
    cor = GRAVIDADE_CORES.get(gravidade, RGBColor(0x00, 0x00, 0x00))

    heading = document.add_paragraph()
    run = heading.add_run(f"{idx}. {ponto.get('titulo', 'Ponto sem título')}")
    run.bold = True
    run.font.size = Pt(12)

    grav_p = document.add_paragraph()
    grav_run = grav_p.add_run(f"Gravidade: {gravidade or 'não informada'}")
    grav_run.bold = True
    grav_run.font.color.rgb = cor
    if ponto.get("categoria"):
        grav_p.add_run(f"   |   Categoria: {ponto['categoria']}")

    if ponto.get("trecho_ou_clausula"):
        p = document.add_paragraph()
        p.add_run("Cláusula/trecho: ").italic = True
        p.add_run(ponto["trecho_ou_clausula"]).italic = True

    if ponto.get("problema"):
        p = document.add_paragraph()
        p.add_run("Problema: ").bold = True
        p.add_run(ponto["problema"])

    if ponto.get("sugestao_de_ajuste"):
        p = document.add_paragraph()
        p.add_run("Sugestão de ajuste: ").bold = True
        p.add_run(ponto["sugestao_de_ajuste"])

    document.add_paragraph()
