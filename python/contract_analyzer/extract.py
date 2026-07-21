"""Extração de texto de contratos em .docx ou .pdf."""

from __future__ import annotations

import os

MIN_USEFUL_CHARS = 200


class ExtractionError(Exception):
    pass


def extract_text(path: str) -> str:
    """Extrai o texto de um arquivo .docx ou .pdf."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".pdf":
        text = _extract_pdf(path)
    else:
        raise ExtractionError(
            f"Extensão '{ext}' não suportada. Use um arquivo .docx ou .pdf."
        )

    if len(text.strip()) < MIN_USEFUL_CHARS:
        raise ExtractionError(
            "Foi extraído muito pouco texto do arquivo. Se o PDF for uma imagem "
            "escaneada (sem texto selecionável), rode um OCR nele antes de analisar."
        )
    return text


def _extract_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_pdf(path: str) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(page_text)
    return "\n".join(parts)
